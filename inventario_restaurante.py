"""
Script para generar un archivo Excel de Control de Inventario y Compras
para restaurantes, usando la librería openpyxl.
"""

from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, GradientFill
)
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from datetime import date
import os

# ---------------------------------------------------------------------------
# PALETA DE COLORES
# ---------------------------------------------------------------------------
COLOR_HEADER_FILL    = "1F3864"   # Azul marino oscuro
COLOR_HEADER_FONT    = "FFFFFF"   # Blanco
COLOR_SUBHEADER_FILL = "2E75B6"   # Azul medio (fila totales / subtítulos)
COLOR_ROW_ODD        = "EAF2FB"   # Azul muy claro (filas impares)
COLOR_ROW_EVEN       = "FFFFFF"   # Blanco (filas pares)
COLOR_ALERT_FILL     = "FFD7D7"   # Rojo claro (SOLICITAR)
COLOR_ALERT_FONT     = "C00000"   # Rojo oscuro (SOLICITAR)
COLOR_OK_FILL        = "D9EAD3"   # Verde claro (OK)
COLOR_OK_FONT        = "1E6B33"   # Verde oscuro (OK)
COLOR_BORDER         = "B8CCE4"   # Borde azul suave
COLOR_DASHBOARD_BG   = "F2F7FD"   # Fondo dashboard
COLOR_TITLE_FILL     = "13274F"   # Azul muy oscuro (título principal)
COLOR_ACCENT         = "F4B942"   # Amarillo dorado (acento)

# ---------------------------------------------------------------------------
# ESTILOS REUTILIZABLES
# ---------------------------------------------------------------------------
def make_border(color=COLOR_BORDER):
    side = Side(style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)

def header_style(ws, cell, text, fill_color=COLOR_HEADER_FILL,
                 font_color=COLOR_HEADER_FONT, size=11, bold=True):
    cell.value = text
    cell.font = Font(name="Calibri", bold=bold, color=font_color, size=size)
    cell.fill = PatternFill("solid", fgColor=fill_color)
    cell.alignment = Alignment(horizontal="center", vertical="center",
                               wrap_text=True)
    cell.border = make_border()

def data_style(cell, fill_color=COLOR_ROW_ODD, font_color="000000",
               align="center", bold=False, size=10):
    cell.font = Font(name="Calibri", color=font_color, size=size, bold=bold)
    cell.fill = PatternFill("solid", fgColor=fill_color)
    cell.alignment = Alignment(horizontal=align, vertical="center",
                               wrap_text=True)
    cell.border = make_border()

# ---------------------------------------------------------------------------
# DATOS DE MUESTRA POR CATEGORÍA
# ---------------------------------------------------------------------------
# Formato: (ID, Producto, Unidad, Stock_Min, Stock_Actual, Proveedor)
SAMPLE_DATA = {
    "Carnes": [
        ("CAR-001", "Pollo entero",          "Kg",      10,  8,  "Proveedora Avícola S.A."),
        ("CAR-002", "Res molida",            "Kg",       5, 12,  "Carnicería Del Valle"),
        ("CAR-003", "Costilla de cerdo",     "Kg",       4,  4,  "Carnicería Del Valle"),
        ("CAR-004", "Filete de res",         "Kg",       3,  1,  "Frigorífico Central"),
        ("CAR-005", "Pechuga de pollo",      "Kg",       8, 15,  "Proveedora Avícola S.A."),
        ("CAR-006", "Tocino ahumado",        "Kg",       2,  2,  "Frigorífico Central"),
        ("CAR-007", "Chorizo español",       "Kg",       3,  7,  "Embutidos Selectos"),
        ("CAR-008", "Camarón mediano",       "Kg",       5,  3,  "Mar Fresco Distribuidora"),
    ],
    "Abarrotes": [
        ("ABA-001", "Arroz blanco",          "Kg",      20, 35,  "Distribuidora La Cosecha"),
        ("ABA-002", "Aceite vegetal",        "Litro",   10,  9,  "Distribuidora La Cosecha"),
        ("ABA-003", "Harina de trigo",       "Kg",      15, 14,  "Molinos del Norte"),
        ("ABA-004", "Sal de mesa",           "Kg",       5, 18,  "Distribuidora La Cosecha"),
        ("ABA-005", "Azúcar estándar",       "Kg",      10,  6,  "Distribuidora La Cosecha"),
        ("ABA-006", "Pasta fideo",           "Kg",       8, 20,  "Distribuidora La Cosecha"),
        ("ABA-007", "Tomate en lata",        "Pieza",   24, 12,  "Conservas Premium"),
        ("ABA-008", "Frijol negro",          "Kg",      10, 22,  "Distribuidora La Cosecha"),
        ("ABA-009", "Consomé de pollo",      "Pieza",   20, 18,  "Condimentos Gourmet"),
        ("ABA-010", "Aceite de oliva",       "Litro",    3,  2,  "Importadora Mediterráneo"),
    ],
    "Desechables": [
        ("DES-001", "Vasos térmicos 8oz",    "Paquete",  5,  3,  "Desechables Express"),
        ("DES-002", "Platos de unicel #3",   "Paquete",  8,  8,  "Desechables Express"),
        ("DES-003", "Tenedores plástico",    "Paquete",  6,  2,  "Desechables Express"),
        ("DES-004", "Servilletas blancas",   "Paquete", 10, 15,  "Papelería Total"),
        ("DES-005", "Bolsas para llevar",    "Paquete", 10,  4,  "Empaques y Más"),
        ("DES-006", "Contenedores 1L",       "Paquete",  5,  5,  "Empaques y Más"),
        ("DES-007", "Cucharas plástico",     "Paquete",  6,  9,  "Desechables Express"),
        ("DES-008", "Papel film rollo",      "Pieza",    3,  1,  "Papelería Total"),
    ],
    "Químicos y Limpieza": [
        ("QUI-001", "Cloro multiusos 1L",    "Litro",    4,  3,  "Limpieza Industrial S.A."),
        ("QUI-002", "Jabón líquido manos",   "Litro",    5,  8,  "Limpieza Industrial S.A."),
        ("QUI-003", "Detergente loza",       "Kg",       4,  2,  "Productos Químicos MX"),
        ("QUI-004", "Sanitizante cocina",    "Litro",    3,  3,  "Productos Químicos MX"),
        ("QUI-005", "Cubetas con tapa",      "Pieza",    2,  5,  "Ferretería Central"),
        ("QUI-006", "Fibras verdes",         "Paquete",  3,  1,  "Limpieza Industrial S.A."),
        ("QUI-007", "Bolsas de basura 80L",  "Paquete",  4,  6,  "Empaques y Más"),
        ("QUI-008", "Guantes de hule M",     "Par",      6,  4,  "Ferretería Central"),
    ],
    "Postres": [
        ("POS-001", "Chocolate amargo 70%",  "Kg",       2,  1,  "Dulcería Gourmet"),
        ("POS-002", "Crema batida litro",    "Litro",    4,  6,  "Lácteos del Campo"),
        ("POS-003", "Flan napolitano",       "Pieza",   10,  8,  "Postres Artesanales"),
        ("POS-004", "Mezcla p/brownies",     "Kg",       3,  2,  "Dulcería Gourmet"),
        ("POS-005", "Chantilly polvo",       "Kg",       2,  3,  "Dulcería Gourmet"),
        ("POS-006", "Helado vainilla 1L",    "Litro",    6,  4,  "Helados Premium"),
        ("POS-007", "Galletas Oreo granel",  "Kg",       2,  2,  "Distribuidora La Cosecha"),
        ("POS-008", "Cajeta de leche",       "Kg",       2,  5,  "Lácteos del Campo"),
    ],
    "Lácteos": [
        ("LAC-001", "Leche entera 1L",       "Litro",   20, 15,  "Lácteos del Campo"),
        ("LAC-002", "Queso Oaxaca",          "Kg",       5,  4,  "Quesería Artesanal"),
        ("LAC-003", "Crema ácida",           "Litro",    6, 10,  "Lácteos del Campo"),
        ("LAC-004", "Mantequilla sin sal",   "Kg",       3,  2,  "Lácteos del Campo"),
        ("LAC-005", "Queso manchego",        "Kg",       4,  4,  "Quesería Artesanal"),
        ("LAC-006", "Yogur natural 1L",      "Litro",    4,  7,  "Lácteos del Campo"),
        ("LAC-007", "Crema para cocinar",    "Litro",    5,  3,  "Lácteos del Campo"),
        ("LAC-008", "Queso crema 1Kg",       "Kg",       3,  1,  "Quesería Artesanal"),
    ],
}

CATEGORY_ICONS = {
    "Carnes":             "🥩",
    "Abarrotes":          "🛒",
    "Desechables":        "🥡",
    "Químicos y Limpieza":"🧴",
    "Postres":            "🍮",
    "Lácteos":            "🥛",
}

COLUMNS = [
    ("ID / Código",            12),
    ("Producto",               32),
    ("Unidad de Medida",       16),
    ("Stock Mínimo",           14),
    ("Stock Actual",           13),
    ("Estado",                 12),
    ("Última Fecha\nde Conteo",18),
    ("Proveedor",              28),
]

TODAY = date.today().strftime("%d/%m/%Y")


# ---------------------------------------------------------------------------
# FUNCIÓN: CREAR PESTAÑA DE CATEGORÍA
# ---------------------------------------------------------------------------
def create_category_sheet(wb: Workbook, category: str, data: list) -> None:
    ws = wb.create_sheet(title=category)
    ws.sheet_view.showGridLines = False

    # ---- FILA 1: Título principal ------------------------------------------
    ws.row_dimensions[1].height = 40
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    icon = CATEGORY_ICONS.get(category, "")
    title_cell.value = f"  {icon}  CONTROL DE INVENTARIO — {category.upper()}"
    title_cell.font = Font(name="Calibri", bold=True, size=16,
                           color=COLOR_HEADER_FONT)
    title_cell.fill = PatternFill("solid", fgColor=COLOR_TITLE_FILL)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    title_cell.border = make_border("13274F")

    # ---- FILA 2: Subtítulo con fecha ---------------------------------------
    ws.row_dimensions[2].height = 18
    ws.merge_cells("A2:H2")
    sub_cell = ws["A2"]
    sub_cell.value = (
        f"  Última actualización: {TODAY}   |   "
        f"Total de productos: {len(data)}"
    )
    sub_cell.font = Font(name="Calibri", italic=True, size=10,
                         color="AAAAAA")
    sub_cell.fill = PatternFill("solid", fgColor="F0F4FA")
    sub_cell.alignment = Alignment(horizontal="left", vertical="center")

    # ---- FILA 3: Encabezados -----------------------------------------------
    ws.row_dimensions[3].height = 38
    for col_idx, (col_name, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=3, column=col_idx)
        header_style(ws, cell, col_name, size=10)

    # ---- FILAS DE DATOS (empiezan en la fila 4) ----------------------------
    data_start_row = 4
    for row_offset, row_data in enumerate(data):
        row_num = data_start_row + row_offset
        ws.row_dimensions[row_num].height = 22

        fill_color = COLOR_ROW_ODD if row_offset % 2 == 0 else COLOR_ROW_EVEN

        # ID
        c = ws.cell(row=row_num, column=1, value=row_data[0])
        data_style(c, fill_color=fill_color, align="center")

        # Producto
        c = ws.cell(row=row_num, column=2, value=row_data[1])
        data_style(c, fill_color=fill_color, align="left")

        # Unidad
        c = ws.cell(row=row_num, column=3, value=row_data[2])
        data_style(c, fill_color=fill_color, align="center")

        # Stock Mínimo
        c = ws.cell(row=row_num, column=4, value=row_data[3])
        data_style(c, fill_color=fill_color, align="center", bold=True)

        # Stock Actual
        c = ws.cell(row=row_num, column=5, value=row_data[4])
        data_style(c, fill_color=fill_color, align="center", bold=True)

        # Estado — fórmula Excel (col E = Stock Actual, col D = Stock Mínimo)
        estado_col = get_column_letter(6)
        stock_actual_col = get_column_letter(5)
        stock_min_col = get_column_letter(4)
        formula = (
            f'=IF({stock_actual_col}{row_num}<={stock_min_col}{row_num},'
            f'"SOLICITAR","OK")'
        )
        c = ws.cell(row=row_num, column=6, value=formula)
        data_style(c, fill_color=fill_color, align="center", bold=True)

        # Última fecha de conteo
        c = ws.cell(row=row_num, column=7, value=TODAY)
        data_style(c, fill_color=fill_color, align="center")

        # Proveedor
        c = ws.cell(row=row_num, column=8, value=row_data[5])
        data_style(c, fill_color=fill_color, align="left")

    # ---- ANCHOS DE COLUMNA -------------------------------------------------
    for col_idx, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ---- FORMATO CONDICIONAL — columna Estado (F) --------------------------
    last_data_row = data_start_row + len(data) - 1
    estado_range = f"F{data_start_row}:F{last_data_row}"

    # SOLICITAR → fondo rojo claro, texto rojo oscuro
    red_fill = PatternFill("solid", fgColor=COLOR_ALERT_FILL)
    red_font = Font(name="Calibri", bold=True, color=COLOR_ALERT_FONT,
                    size=10)
    ws.conditional_formatting.add(
        estado_range,
        FormulaRule(
            formula=[f'F{data_start_row}="SOLICITAR"'],
            fill=red_fill,
            font=red_font,
        ),
    )

    # OK → fondo verde claro, texto verde oscuro
    green_fill = PatternFill("solid", fgColor=COLOR_OK_FILL)
    green_font = Font(name="Calibri", bold=True, color=COLOR_OK_FONT,
                      size=10)
    ws.conditional_formatting.add(
        estado_range,
        FormulaRule(
            formula=[f'F{data_start_row}="OK"'],
            fill=green_fill,
            font=green_font,
        ),
    )

    # ---- FILA VACÍA DE CIERRE (separación visual) -------------------------
    ws.row_dimensions[last_data_row + 1].height = 10

    # ---- FREEZE PANES (congela encabezados) --------------------------------
    ws.freeze_panes = ws["A4"]

    # ---- ZONA DE IMPRESIÓN -------------------------------------------------
    ws.print_area = f"A1:H{last_data_row}"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.orientation = "landscape"


# ---------------------------------------------------------------------------
# FUNCIÓN: CREAR DASHBOARD
# ---------------------------------------------------------------------------
def create_dashboard(wb: Workbook, all_data: dict) -> None:
    ws = wb.create_sheet(title="DASHBOARD", index=0)
    ws.sheet_view.showGridLines = False

    # ---- TÍTULO PRINCIPAL --------------------------------------------------
    ws.row_dimensions[1].height = 50
    ws.merge_cells("A1:G1")
    t = ws["A1"]
    t.value = "  🍽️  SISTEMA DE CONTROL DE INVENTARIO Y COMPRAS"
    t.font = Font(name="Calibri", bold=True, size=20,
                  color=COLOR_HEADER_FONT)
    t.fill = PatternFill("solid", fgColor=COLOR_TITLE_FILL)
    t.alignment = Alignment(horizontal="left", vertical="center")

    # ---- SUBTÍTULO ---------------------------------------------------------
    ws.row_dimensions[2].height = 20
    ws.merge_cells("A2:G2")
    s = ws["A2"]
    s.value = f"  Dashboard de alertas — Generado: {TODAY}"
    s.font = Font(name="Calibri", italic=True, size=11, color="888888")
    s.fill = PatternFill("solid", fgColor="F0F4FA")
    s.alignment = Alignment(horizontal="left", vertical="center")

    # ---- SEPARADOR ---------------------------------------------------------
    ws.row_dimensions[3].height = 8
    ws.merge_cells("A3:G3")
    ws["A3"].fill = PatternFill("solid", fgColor=COLOR_ACCENT)

    # ---- TARJETAS DE RESUMEN POR CATEGORÍA ---------------------------------
    ws.row_dimensions[4].height = 30
    ws.merge_cells("A4:G4")
    sec = ws["A4"]
    sec.value = "  📊  RESUMEN GENERAL POR CATEGORÍA"
    sec.font = Font(name="Calibri", bold=True, size=12,
                    color=COLOR_HEADER_FONT)
    sec.fill = PatternFill("solid", fgColor=COLOR_SUBHEADER_FILL)
    sec.alignment = Alignment(horizontal="left", vertical="center")

    # Encabezados de resumen
    ws.row_dimensions[5].height = 30
    summary_headers = [
        "Categoría", "Total\nProductos", "En Stock OK",
        "Stock Crítico", "% Alertas", "Acción Requerida", ""
    ]
    summary_widths = [24, 14, 13, 14, 12, 22, 4]

    for col_idx, (hdr, w) in enumerate(
            zip(summary_headers, summary_widths), start=1):
        cell = ws.cell(row=5, column=col_idx)
        header_style(ws, cell, hdr, size=10)
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    # Filas de resumen
    current_row = 6
    total_criticos_global = 0
    for cat_idx, (category, data) in enumerate(all_data.items()):
        ws.row_dimensions[current_row].height = 22
        icon = CATEGORY_ICONS.get(category, "")
        total = len(data)
        criticos = sum(1 for d in data if d[4] <= d[3])   # stock_actual <= stock_min
        ok_count = total - criticos
        pct = round(criticos / total * 100) if total else 0
        total_criticos_global += criticos

        fill_c = COLOR_ROW_ODD if cat_idx % 2 == 0 else COLOR_ROW_EVEN

        vals = [
            f" {icon}  {category}",
            total,
            ok_count,
            criticos,
            f"{pct}%",
            "⚠️  REVISAR PENDIENTES" if criticos > 0 else "✅  TODO EN ORDEN",
            "",
        ]
        for col_idx, val in enumerate(vals, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            align = "left" if col_idx in (1, 6) else "center"
            bold = col_idx in (4, 6)
            font_color = (COLOR_ALERT_FONT if (col_idx == 4 and criticos > 0)
                          else COLOR_OK_FONT if col_idx == 4
                          else "000000")
            row_fill = (COLOR_ALERT_FILL if (col_idx == 4 and criticos > 0)
                        else COLOR_OK_FILL if col_idx == 4
                        else fill_c)
            data_style(cell, fill_color=row_fill, font_color=font_color,
                       align=align, bold=bold)
        current_row += 1

    # ---- SEPARADOR ---------------------------------------------------------
    current_row += 1
    ws.row_dimensions[current_row].height = 8
    ws.merge_cells(f"A{current_row}:G{current_row}")
    ws[f"A{current_row}"].fill = PatternFill("solid", fgColor=COLOR_ACCENT)
    current_row += 1

    # ---- SECCIÓN: LISTA DE STOCK CRÍTICO -----------------------------------
    ws.row_dimensions[current_row].height = 30
    ws.merge_cells(f"A{current_row}:G{current_row}")
    sec2 = ws[f"A{current_row}"]
    sec2.value = "  🚨  STOCK CRÍTICO — PRODUCTOS QUE NECESITAN PEDIRSE YA"
    sec2.font = Font(name="Calibri", bold=True, size=12,
                     color=COLOR_HEADER_FONT)
    sec2.fill = PatternFill("solid", fgColor="C00000")
    sec2.alignment = Alignment(horizontal="left", vertical="center")
    current_row += 1

    # Encabezados tabla críticos
    ws.row_dimensions[current_row].height = 30
    crit_headers = [
        "Categoría", "ID / Código", "Producto",
        "Stock Mínimo", "Stock Actual", "Déficit", "Proveedor"
    ]
    for col_idx, hdr in enumerate(crit_headers, start=1):
        cell = ws.cell(row=current_row, column=col_idx)
        header_style(ws, cell, hdr, fill_color="7F0000", size=10)
    current_row += 1

    # Filas críticas
    alert_row_idx = 0
    for category, data in all_data.items():
        icon = CATEGORY_ICONS.get(category, "")
        for row_data in data:
            id_, product, unit, stock_min, stock_actual, supplier = row_data
            if stock_actual <= stock_min:
                ws.row_dimensions[current_row].height = 20
                deficit = stock_min - stock_actual
                fill_c = COLOR_ALERT_FILL if alert_row_idx % 2 == 0 else "FFE5E5"
                alert_vals = [
                    f"{icon}  {category}", id_, product,
                    stock_min, stock_actual, deficit, supplier
                ]
                for col_idx, val in enumerate(alert_vals, start=1):
                    cell = ws.cell(row=current_row, column=col_idx, value=val)
                    align = "left" if col_idx in (1, 3, 7) else "center"
                    bold = col_idx in (4, 5, 6)
                    font_color = COLOR_ALERT_FONT if col_idx == 6 else "000000"
                    data_style(cell, fill_color=fill_c,
                               font_color=font_color,
                               align=align, bold=bold)
                current_row += 1
                alert_row_idx += 1

    if alert_row_idx == 0:
        ws.row_dimensions[current_row].height = 26
        ws.merge_cells(f"A{current_row}:G{current_row}")
        ok_cell = ws[f"A{current_row}"]
        ok_cell.value = "   ✅  No hay productos en stock crítico en este momento."
        ok_cell.font = Font(name="Calibri", size=11,
                            color=COLOR_OK_FONT, bold=True)
        ok_cell.fill = PatternFill("solid", fgColor=COLOR_OK_FILL)
        ok_cell.alignment = Alignment(horizontal="left", vertical="center")
        current_row += 1

    # ---- FOOTER ------------------------------------------------------------
    current_row += 1
    ws.row_dimensions[current_row].height = 20
    ws.merge_cells(f"A{current_row}:G{current_row}")
    footer = ws[f"A{current_row}"]
    footer.value = (
        "  ℹ️  Los datos de inventario se actualizan manualmente en cada "
        "pestaña de categoría.  |  Navega a cada pestaña para registrar "
        "conteos físicos."
    )
    footer.font = Font(name="Calibri", italic=True, size=9, color="888888")
    footer.fill = PatternFill("solid", fgColor="F0F4FA")
    footer.alignment = Alignment(horizontal="left", vertical="center")

    ws.freeze_panes = ws["A6"]


# ---------------------------------------------------------------------------
# FUNCIÓN PRINCIPAL
# ---------------------------------------------------------------------------
def build_workbook(output_path: str = "Inventario_Restaurante.xlsx") -> str:
    wb = Workbook()
    # Eliminar la hoja por defecto
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    # Crear Dashboard primero
    create_dashboard(wb, SAMPLE_DATA)

    # Crear pestaña por categoría
    for category, data in SAMPLE_DATA.items():
        create_category_sheet(wb, category, data)

    # Asignar colores de pestaña (tab color)
    tab_colors = {
        "DASHBOARD":          "13274F",
        "Carnes":             "C00000",
        "Abarrotes":          "375623",
        "Desechables":        "7030A0",
        "Químicos y Limpieza":"1F3864",
        "Postres":            "C55A11",
        "Lácteos":            "2E75B6",
    }
    for sheet_name, color in tab_colors.items():
        if sheet_name in wb.sheetnames:
            wb[sheet_name].sheet_properties.tabColor = color

    wb.save(output_path)
    return os.path.abspath(output_path)


# ---------------------------------------------------------------------------
# ENTRY POINT
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    path = build_workbook()
    print(f"\n✅  Archivo generado exitosamente:")
    print(f"   {path}\n")
    print("📋  Pestañas creadas:")
    print("    • DASHBOARD    — Vista general con alertas de stock crítico")
    for cat, icon in CATEGORY_ICONS.items():
        n = len(SAMPLE_DATA[cat])
        criticos = sum(1 for d in SAMPLE_DATA[cat] if d[4] <= d[3])
        alerta = f"  ⚠️  {criticos} en STOCK CRÍTICO" if criticos else "  ✅  Todo en orden"
        print(f"    • {icon}  {cat:<22} ({n} productos){alerta}")
    print()
