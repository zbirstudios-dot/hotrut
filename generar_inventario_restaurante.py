from __future__ import annotations

import argparse
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties


COLUMN_HEADERS = [
    "ID / Código",
    "Producto (Nombre)",
    "Unidad de Medida (Kg, Pieza, Litro, Paquete)",
    "Stock Mínimo (El límite antes de alarmar)",
    "Stock Actual (Para el conteo físico)",
    "Estado",
    "Última Fecha de Conteo",
    "Proveedor",
]

CATEGORIES = [
    {"display_name": "Carnes", "sheet_name": "Carnes"},
    {"display_name": "Abarrotes", "sheet_name": "Abarrotes"},
    {"display_name": "Desechables", "sheet_name": "Desechables"},
    {
        "display_name": "Químicos / Limpieza",
        "sheet_name": "Químicos - Limpieza",
    },
    {"display_name": "Postres", "sheet_name": "Postres"},
    {"display_name": "Lácteos", "sheet_name": "Lácteos"},
]

DATA_ROW_COUNT = 100
FIRST_DATA_ROW = 2
LAST_DATA_ROW = FIRST_DATA_ROW + DATA_ROW_COUNT - 1

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

TITLE_FILL = PatternFill("solid", fgColor="DCE6F1")
TITLE_FONT = Font(color="1F1F1F", bold=True, size=14)
SUBTITLE_FONT = Font(color="4F4F4F", italic=True)

BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

LIGHT_ROW_FILL = PatternFill("solid", fgColor="F7F9FC")
CRITICAL_FILL = PatternFill("solid", fgColor="FCE4D6")
CRITICAL_FONT = Font(color="9C0006", bold=True)


def apply_header_style(cell) -> None:
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = HEADER_ALIGNMENT
    cell.border = BORDER


def apply_body_style(cell, centered: bool = False) -> None:
    cell.border = BORDER
    cell.alignment = Alignment(
        horizontal="center" if centered else "left",
        vertical="center",
        wrap_text=True,
    )


def build_status_formula(row_number: int) -> str:
    return f'=IF(OR(B{row_number}="",D{row_number}="",E{row_number}=""),"",IF(E{row_number}<=D{row_number},"SOLICITAR","OK"))'


def setup_category_sheet(ws, config: dict[str, str]) -> None:
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = f"A1:H{LAST_DATA_ROW}"

    for column_index, header in enumerate(COLUMN_HEADERS, start=1):
        header_cell = ws.cell(row=1, column=column_index, value=header)
        apply_header_style(header_cell)

    for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        for column in range(1, 9):
            cell = ws.cell(row=row, column=column)
            centered_columns = {1, 3, 4, 5, 6, 7}
            apply_body_style(cell, centered=column in centered_columns)

            if row % 2 == 0:
                cell.fill = LIGHT_ROW_FILL

        ws[f"F{row}"] = build_status_formula(row)
        ws[f"G{row}"].number_format = "DD/MM/YYYY"

    critical_rule = FormulaRule(
        formula=[f'EXACT($F{FIRST_DATA_ROW},"SOLICITAR")'],
        fill=CRITICAL_FILL,
        font=CRITICAL_FONT,
    )
    ws.conditional_formatting.add(
        f"F{FIRST_DATA_ROW}:F{LAST_DATA_ROW}",
        critical_rule,
    )

    width_overrides = {
        "A": 14,
        "B": 28,
        "C": 22,
        "D": 18,
        "E": 16,
        "F": 14,
        "G": 18,
        "H": 24,
    }
    auto_adjust_widths(ws, width_overrides, max_width=32)


def build_dashboard_detail_formula(sheet_name: str, row_number: int) -> str:
    quoted_sheet = f"'{sheet_name}'"
    return (
        f'=IF(B{row_number}=0,"Sin alertas",'
        f'TEXTJOIN(CHAR(10),TRUE,'
        f'IF({quoted_sheet}!$F$2:$F${LAST_DATA_ROW}="SOLICITAR",'
        f'{quoted_sheet}!$B$2:$B${LAST_DATA_ROW}'
        f'&" | Stock: "&{quoted_sheet}!$E$2:$E${LAST_DATA_ROW}'
        f'&" | Mín.: "&{quoted_sheet}!$D$2:$D${LAST_DATA_ROW}'
        f'&IF({quoted_sheet}!$H$2:$H${LAST_DATA_ROW}<>""," | Prov.: "&{quoted_sheet}!$H$2:$H${LAST_DATA_ROW},""),'
        f'"")))'
    )


def setup_dashboard_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A6"

    ws.merge_cells("A1:C1")
    ws["A1"] = "DASHBOARD - Control de Inventario y Compras"
    ws["A1"].fill = TITLE_FILL
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:C2")
    ws["A2"] = (
        "Panel de alertas automáticas para identificar productos en stock crítico "
        "y priorizar compras."
    )
    ws["A2"].font = SUBTITLE_FONT
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws["A4"] = "Fecha de generación"
    ws["B4"] = datetime.now().strftime("%d/%m/%Y %H:%M")
    ws["A4"].font = Font(bold=True)

    ws["A5"] = "Departamento"
    ws["B5"] = "Alertas críticas"
    ws["C5"] = "Productos por solicitar"

    for cell in ws[5]:
        apply_header_style(cell)

    for index, category in enumerate(CATEGORIES, start=6):
        quoted_sheet = f"'{category['sheet_name']}'"
        ws[f"A{index}"] = category["display_name"]
        ws[f"B{index}"] = f'=COUNTIF({quoted_sheet}!$F$2:$F${LAST_DATA_ROW},"SOLICITAR")'
        ws[f"C{index}"] = build_dashboard_detail_formula(category["sheet_name"], index)

        for column in ("A", "B", "C"):
            cell = ws[f"{column}{index}"]
            cell.border = BORDER
            cell.alignment = Alignment(
                horizontal="center" if column == "B" else "left",
                vertical="top",
                wrap_text=True,
            )

    ws["A13"] = "Total de alertas críticas"
    ws["B13"] = "=SUM(B6:B11)"
    ws["A13"].font = Font(bold=True)
    ws["B13"].font = Font(bold=True, color="9C0006")

    ws.conditional_formatting.add(
        "B6:B11",
        FormulaRule(
            formula=["B6>0"],
            fill=CRITICAL_FILL,
            font=CRITICAL_FONT,
        ),
    )

    for row in range(6, 12):
        ws.row_dimensions[row].height = 54

    auto_adjust_widths(
        ws,
        {
            "A": 28,
            "B": 16,
            "C": 65,
        },
        max_width=65,
    )


def auto_adjust_widths(ws, minimum_widths: dict[str, int] | None = None, max_width: int = 60) -> None:
    minimum_widths = minimum_widths or {}

    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = minimum_widths.get(column_letter, 10)

        for cell in column_cells:
            if cell.value is None:
                continue

            if cell.data_type == "f":
                if column_letter == "F":
                    value_length = len("SOLICITAR")
                elif column_letter == "C" and ws.title == "DASHBOARD":
                    value_length = minimum_widths.get(column_letter, 60)
                else:
                    value_length = len(str(cell.value))
            else:
                value_length = len(str(cell.value))

            max_length = max(max_length, value_length + 2)

        ws.column_dimensions[column_letter].width = min(max_length, max_width)


def create_inventory_workbook(output_path: str) -> str:
    workbook = Workbook()
    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True)

    dashboard_sheet = workbook.active
    dashboard_sheet.title = "DASHBOARD"
    setup_dashboard_sheet(dashboard_sheet)

    for category in CATEGORIES:
        category_sheet = workbook.create_sheet(title=category["sheet_name"])
        setup_category_sheet(category_sheet, category)

    workbook.save(output_path)
    return output_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Genera un archivo Excel para control de inventario y compras."
    )
    parser.add_argument(
        "-o",
        "--output",
        default="inventario_restaurante.xlsx",
        help="Ruta del archivo Excel a generar.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output_path = create_inventory_workbook(args.output)
    print(f"Archivo generado correctamente: {output_path}")
    print("Nota: la pestaña 'Químicos / Limpieza' se guarda como 'Químicos - Limpieza' porque Excel no permite '/'.")


if __name__ == "__main__":
    main()
